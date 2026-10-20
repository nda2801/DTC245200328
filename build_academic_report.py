import os
import re
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_academic_docx():
    workspace_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(workspace_dir, "docs", "Bao_Cao_Trien_Khai_He_Thong.md")
    docx_path = os.path.join(workspace_dir, "docs", "Bao_Cao_Trien_Khai_He_Thong.docx")
    root_docx_path = os.path.join(workspace_dir, "Bao_Cao_Trien_Khai_He_Thong.docx")
    images_dir = os.path.join(workspace_dir, "docs", "images")
    diagrams_dir = os.path.join(images_dir, "diagrams")

    if not os.path.exists(md_path):
        print(f"Error: Markdown file not found at {md_path}")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = docx.Document()

    # 1. Automatic Field Update in Microsoft Word
    settings = doc.settings.element
    update_fields = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
    settings.append(update_fields)

    # 2. Page Geometry: Chuẩn Thể thức Văn bản Học thuật Việt Nam (Nghị định 30/2020/NĐ-CP)
    # Khổ A4: Lề Trên 20mm, Dưới 20mm, Trái 30mm (đóng gáy), Phải 20mm
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(20)

    # Header - 100% Màu Đen Thuần Túy
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Báo cáo Kết thúc Học phần: Triển khai và Quản trị Hệ thống Phần mềm - Đề tài 2 (ICTU)")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(0, 0, 0)

    # Footer - 100% Màu Đen Thuần Túy
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Sinh viên: Nguyễn Đức Anh - Mã SV: DTC245200328 - Lớp: CNTT - K23B")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(9)
    frun.font.color.rgb = RGBColor(0, 0, 0)

    # 3. Configure Built-in Word Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.3
    normal_style.paragraph_format.space_after = Pt(4)

    h1_style = doc.styles['Heading 1']
    h1_style.font.name = 'Times New Roman'
    h1_style.font.size = Pt(18)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0, 0, 0)
    h1_style.paragraph_format.space_before = Pt(14)
    h1_style.paragraph_format.space_after = Pt(6)
    h1_style.paragraph_format.keep_with_next = True

    h2_style = doc.styles['Heading 2']
    h2_style.font.name = 'Times New Roman'
    h2_style.font.size = Pt(17)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0, 0, 0)
    h2_style.paragraph_format.space_before = Pt(10)
    h2_style.paragraph_format.space_after = Pt(4)
    h2_style.paragraph_format.keep_with_next = True

    h3_style = doc.styles['Heading 3']
    h3_style.font.name = 'Times New Roman'
    h3_style.font.size = Pt(16)
    h3_style.font.bold = True
    h3_style.font.color.rgb = RGBColor(0, 0, 0)
    h3_style.paragraph_format.space_before = Pt(8)
    h3_style.paragraph_format.space_after = Pt(2)
    h3_style.paragraph_format.keep_with_next = True

    try:
        h4_style = doc.styles['Heading 4']
    except KeyError:
        h4_style = doc.styles.add_style('Heading 4', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    h4_style.font.name = 'Times New Roman'
    h4_style.font.size = Pt(14)
    h4_style.font.bold = True
    h4_style.font.color.rgb = RGBColor(0, 0, 0)
    h4_style.paragraph_format.space_before = Pt(6)
    h4_style.paragraph_format.space_after = Pt(2)
    h4_style.paragraph_format.keep_with_next = True

    try:
        caption_style = doc.styles['Caption']
    except KeyError:
        caption_style = doc.styles.add_style('Caption', docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    caption_style.font.name = 'Times New Roman'
    caption_style.font.size = Pt(11)
    caption_style.font.italic = True
    caption_style.font.color.rgb = RGBColor(0, 0, 0)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_style.paragraph_format.space_before = Pt(2)
    caption_style.paragraph_format.space_after = Pt(8)

    # Helper XML functions
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>\n'
            f'  <w:top w:w="{top}" w:type="dxa"/>\n'
            f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
            f'  <w:left w:w="{left}" w:type="dxa"/>\n'
            f'  <w:right w:w="{right}" w:type="dxa"/>\n'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_table_borders(table, color="000000", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_formatted_runs(paragraph, text, base_font_size=13, font_name="Times New Roman", default_color=RGBColor(0, 0, 0)):
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
        for token in tokens:
            if not token:
                continue
            if token.startswith('**') and token.endswith('**') and len(token) >= 4:
                run = paragraph.add_run(token[2:-2])
                run.bold = True
                run.font.name = font_name
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = default_color
            elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
                run = paragraph.add_run(token[1:-1])
                run.italic = True
                run.font.name = font_name
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = default_color
            elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(base_font_size - 1)
                run.font.color.rgb = default_color
            else:
                run = paragraph.add_run(token)
                run.font.name = font_name
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = default_color

    def add_heading_with_outline(doc, text, level):
        style_name = f'Heading {level}'
        p = doc.add_paragraph(style=style_name)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        
        pPr = p._p.get_or_add_pPr()
        outline_lvl = parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level - 1}"/>')
        pPr.append(outline_lvl)
        
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
        if level == 1:
            run.font.size = Pt(18)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif level == 2:
            run.font.size = Pt(17)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif level == 3:
            run.font.size = Pt(16)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif level == 4:
            run.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
        return p

    def add_toc_entry_line(doc, title, page_str, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(160), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(2)
        
        if level == 1:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.left_indent = Inches(0)
            r_title = p.add_run(title)
            r_title.bold = True
            r_title.font.name = "Times New Roman"
            r_title.font.size = Pt(12)
            r_title.font.color.rgb = RGBColor(0, 0, 0)
        elif level == 2:
            p.paragraph_format.left_indent = Inches(0.25)
            r_title = p.add_run(title)
            r_title.font.name = "Times New Roman"
            r_title.font.size = Pt(11.5)
            r_title.font.color.rgb = RGBColor(0, 0, 0)
        else:
            p.paragraph_format.left_indent = Inches(0.45)
            r_title = p.add_run(title)
            r_title.font.name = "Times New Roman"
            r_title.font.size = Pt(11)
            r_title.font.color.rgb = RGBColor(0, 0, 0)
            
        r_tab = p.add_run("\t")
        r_tab.font.name = "Times New Roman"
        r_tab.font.color.rgb = RGBColor(0, 0, 0)
        
        r_page = p.add_run(str(page_str))
        r_page.bold = (level == 1)
        r_page.font.name = "Times New Roman"
        r_page.font.size = Pt(11.5 if level == 1 else 11)
        r_page.font.color.rgb = RGBColor(0, 0, 0)
        return p

    # -------------------------------------------------------------
    # 4. COVER PAGE (Trang bìa đôi viền đen 100% chuẩn ICTU)
    # -------------------------------------------------------------
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_cell = cover_table.rows[0].cells[0]
    cover_cell.width = Mm(160)
    set_cell_margins(cover_cell, top=250, bottom=250, left=350, right=350)

    tcPr = cover_cell._tc.get_or_add_tcPr()
    border_xml = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="double" w:sz="18" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="double" w:sz="18" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="double" w:sz="18" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="double" w:sz="18" w:space="0" w:color="000000"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(border_xml)

    cp = cover_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(2)
    run = cp.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — ĐẠI HỌC THÁI NGUYÊN")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)

    cp2 = cover_cell.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2.paragraph_format.space_after = Pt(2)
    run2 = cp2.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG (ICTU)")
    run2.bold = True
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(12)
    run2.font.color.rgb = RGBColor(0, 0, 0)

    cp3 = cover_cell.add_paragraph()
    cp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp3.paragraph_format.space_after = Pt(8)
    run3 = cp3.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n———————————————")
    run3.bold = True
    run3.font.name = "Times New Roman"
    run3.font.size = Pt(12)
    run3.font.color.rgb = RGBColor(0, 0, 0)

    # Thêm Logo ICTU vào bìa nếu có
    logo_file = os.path.join(diagrams_dir, "ictu_logo.png")
    if os.path.exists(logo_file):
        cp_logo = cover_cell.add_paragraph()
        cp_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp_logo.paragraph_format.space_after = Pt(8)
        cp_logo.add_run().add_picture(logo_file, width=Inches(1.2))

    cp_title1 = cover_cell.add_paragraph()
    cp_title1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_title1.paragraph_format.space_before = Pt(8)
    cp_title1.paragraph_format.space_after = Pt(4)
    run_t1 = cp_title1.add_run("BÁO CÁO KẾT THÚC HỌC PHẦN")
    run_t1.bold = True
    run_t1.font.name = "Times New Roman"
    run_t1.font.size = Pt(16)
    run_t1.font.color.rgb = RGBColor(0, 0, 0)

    cp_title2 = cover_cell.add_paragraph()
    cp_title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_title2.paragraph_format.space_after = Pt(14)
    run_t2 = cp_title2.add_run("HỌC PHẦN: TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    run_t2.bold = True
    run_t2.font.name = "Times New Roman"
    run_t2.font.size = Pt(13)
    run_t2.font.color.rgb = RGBColor(0, 0, 0)

    cp_sub = cover_cell.add_paragraph()
    cp_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_sub.paragraph_format.space_after = Pt(14)
    run_s = cp_sub.add_run(
        "ĐỀ TÀI 2: XÂY DỰNG, TRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG\n"
        "QUẢN LÝ SINH VIÊN TRÊN NỀN TẢNG CONTAINER HÓA\n"
        "VÀ OBSERVABILITY STACK"
    )
    run_s.bold = True
    run_s.font.name = "Times New Roman"
    run_s.font.size = Pt(14)
    run_s.font.color.rgb = RGBColor(0, 0, 0)

    info_table = cover_cell.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_table, color="FFFFFF", sz="0", val="none")

    info_data = [
        ("Sinh viên thực hiện:", "NGUYỄN ĐỨC ANH"),
        ("Mã số sinh viên (MSSV):", "DTC245200328"),
        ("Lớp chuyên ngành:", "CNTT - K23B"),
        ("Ngành đào tạo:", "Công nghệ Thông tin"),
        ("Giảng viên hướng dẫn:", "Bộ môn Mạng & Hệ thống Phần mềm"),
        ("Kho lưu trữ GitHub:", "https://github.com/nda2801/DTC245200328")
    ]
    for idx, (label, val) in enumerate(info_data):
        row = info_table.rows[idx]
        p_lbl = row.cells[0].paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(3)
        r_lbl = p_lbl.add_run(label)
        r_lbl.bold = True
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(12)
        r_lbl.font.color.rgb = RGBColor(0, 0, 0)
        row.cells[0].width = Mm(65)

        p_val = row.cells[1].paragraphs[0]
        p_val.paragraph_format.space_after = Pt(3)
        r_val = p_val.add_run(val)
        if idx in (0, 1, 2, 5):
            r_val.bold = True
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(12)
        r_val.font.color.rgb = RGBColor(0, 0, 0)
        row.cells[1].width = Mm(85)

    cover_cell.add_paragraph()

    cp_loc = cover_cell.add_paragraph()
    cp_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_loc.paragraph_format.space_after = Pt(0)
    r_loc = cp_loc.add_run("THÁI NGUYÊN, THÁNG 10 NĂM 2026")
    r_loc.bold = True
    r_loc.font.name = "Times New Roman"
    r_loc.font.size = Pt(12)
    r_loc.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 5. DỮ LIỆU DANH MỤC CHO 3 HỆ THỐNG MỤC LỤC TỰ ĐỘNG
    # -------------------------------------------------------------
    toc_catalog = [
        ("LỜI NÓI ĐẦU", "2", 1),
        ("1. Lời cam đoan", "2", 2),
        ("2. Lời mở đầu", "2", 2),
        ("MỤC LỤC 1: MỤC LỤC CÁC CHƯƠNG (CHAPTER TOC)", "3", 1),
        ("MỤC LỤC 2: MỤC LỤC CHI TIẾT NỘI DUNG (DETAILED TOC)", "4", 1),
        ("MỤC LỤC 3: DANH MỤC HÌNH ẢNH VÀ BẢNG BIỂU", "5", 1),
        ("DANH MỤC TỪ VIẾT TẮT (ACRONYMS)", "6", 1),
        ("CHƯƠNG 1: CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ ÁP DỤNG", "7", 1),
        ("1.1. Công nghệ Container hóa và Docker, Docker Compose", "7", 2),
        ("1.2. Nền tảng Web Application Node.js Express và CSDL MySQL 8.0, phpMyAdmin", "8", 2),
        ("1.3. Cổng Ingress Gateway Nginx, Giao thức HTTPS và 6 Security Headers", "8", 2),
        ("1.4. Hệ thống Giám sát chuỗi thời gian Prometheus và Grafana (RED & USE)", "9", 2),
        ("1.5. Hệ thống Quản lý nhật ký tập trung Grafana Loki, Promtail và Ngôn ngữ LogQL", "10", 2),
        ("1.6. Các nguyên lý An toàn thông tin và Gia cố hệ thống (Hardening)", "11", 2),
        ("CHƯƠNG 2: KIẾN TRÚC HỆ THỐNG VÀ QUẢN LÝ MÃ NGUỒN GITHUB", "12", 1),
        ("2.1. Mục tiêu và bối cảnh đề tài", "12", 2),
        ("2.2. Mô hình kiến trúc 10 Containers và Phân vùng mạng 3 tầng cô lập", "13", 2),
        ("2.3. Quản lý mã nguồn trên GitHub và Quy chuẩn 3 Commits", "14", 2),
        ("CHƯƠNG 3: TRIỂN KHAI DỊCH VỤ CỐT LÕI VÀ INGRESS GATEWAY (COMMIT 1)", "15", 1),
        ("3.1. Thiết kế cơ sở dữ liệu MySQL (student_db) và Phân quyền Least Privilege", "15", 2),
        ("3.2. Xây dựng dịch vụ Web Application (Node.js Express & Winston Logger)", "16", 2),
        ("3.3. Phát triển giao diện người dùng Enterprise SaaS Dashboard", "17", 2),
        ("3.4. Triển khai công cụ quản trị phpMyAdmin qua đường dẫn bảo vệ /pma/", "18", 2),
        ("3.5. Cấu hình Nginx Ingress Reverse Proxy, SSL tự ký và 6 Security Headers", "19", 2),
        ("CHƯƠNG 4: HỆ THỐNG GIÁM SÁT OBSERVABILITY PROMETHEUS VÀ GRAFANA (COMMIT 2)", "20", 1),
        ("4.1. Thiết lập hạ tầng thu thập chỉ số Prometheus với 4 Scrape Jobs", "20", 2),
        ("4.2. Thực thi kiểm chứng truy vấn PromQL trên Prometheus Graph", "21", 2),
        ("4.3. Thiết kế Dashboard Grafana kết hợp mô hình RED và USE", "22", 2),
        ("4.4. Giám sát chi tiết tài nguyên container (cAdvisor) và hiệu năng MySQL", "23", 2),
        ("4.5. Xử lý đồng bộ Data Source UID và Xuất tệp cấu hình độc lập", "24", 2),
        ("CHƯƠNG 5: QUẢN LÝ LOG TẬP TRUNG LOKI, TRUY VẤN LOGQL VÀ HARDENING (COMMIT 3)", "25", 1),
        ("5.1. Kiến trúc thu thập nhật ký tập trung Grafana Loki & Promtail", "25", 2),
        ("5.2. Xây dựng và kiểm chứng 3 câu truy vấn LogQL chuyên sâu", "26", 2),
        ("5.3. Triển khai chi tiết 5 giải pháp Hardening an toàn hệ thống", "27", 2),
        ("5.4. Xác thực thực nghiệm các tiêu chuẩn Hardening qua câu lệnh CLI", "28", 2),
        ("KẾT LUẬN", "29", 1),
        ("TÀI LIỆU THAM KHẢO", "30", 1),
        ("PHỤ LỤC", "31", 1)
    ]

    figures_catalog = [
        ("Hình 2.1: Sơ đồ kiến trúc vi dịch vụ 10 containers và phân vùng mạng nội bộ", "13"),
        ("Hình 2.2: Minh chứng lịch sử 3 commits đạt chuẩn trên GitHub Repository", "14"),
        ("Hình 3.0: Lược đồ quan hệ thực thể ERD cơ sở dữ liệu student_db", "15"),
        ("Hình 3.1: Giao diện Tổng quan Dashboard Quản lý Sinh viên qua HTTPS", "17"),
        ("Hình 3.2: Bảng dữ liệu danh sách sinh viên ICTU và các chức năng tác vụ", "17"),
        ("Hình 3.3: Modal biểu mẫu thêm mới thông tin hồ sơ sinh viên", "18"),
        ("Hình 3.4: Thông báo Toast phản hồi thời gian thực khi thêm sinh viên thành công", "18"),
        ("Hình 3.5: Giao diện quản lý bảng điểm học phần và tự động quy đổi GPA", "18"),
        ("Hình 3.6: Giao diện quản lý danh mục lớp học chuyên ngành", "19"),
        ("Hình 3.7: Thao tác lọc sinh viên động theo lớp trực thuộc", "19"),
        ("Hình 3.8: Giao diện quản trị cơ sở dữ liệu MySQL qua phpMyAdmin", "19"),
        ("Hình 3.9: Sơ đồ luồng Ingress Nginx SSL Reverse Proxy và định tuyến dịch vụ", "20"),
        ("Hình 4.0: Sơ đồ luồng thu thập metrics Prometheus và trực quan hóa Grafana", "21"),
        ("Hình 4.1: Trạng thái UP đồng thời của 4 Scrape Targets trên Prometheus UI", "21"),
        ("Hình 4.2: Giao diện Prometheus Graph thực thi truy vấn PromQL kiểm tra trạng thái CSDL", "22"),
        ("Hình 4.3: Dashboard giám sát tổng thể thời gian thực kết hợp RED & USE trên Grafana", "23"),
        ("Hình 4.4: Phân hệ giám sát tài nguyên phần cứng cAdvisor và hiệu năng MySQL trên Grafana", "23"),
        ("Hình 5.0: Sơ đồ luồng thu thập và truy vấn nhật ký tập trung Promtail - Loki - Grafana", "25"),
        ("Hình 5.1: Giao diện Grafana Explore thực thi truy vấn LogQL luồng log container từ Loki Server", "26"),
        ("Hình 5.2: Giao diện Grafana Explore lọc nhật ký kiểm toán an ninh [AUDIT] của Web App", "27"),
        ("Hình 5.3: Minh chứng kiểm tra thực nghiệm các giải pháp Hardening an toàn hệ thống", "28")
    ]

    tables_catalog = [
        ("Bảng 2.1: Danh mục 10 containers dịch vụ trong hạ tầng Docker Compose", "13"),
        ("Bảng 2.2: Lịch sử 3 commits bắt buộc trên GitHub theo chuẩn Conventional Commits", "14"),
        ("Bảng 3.1: Lược đồ quan hệ và ràng buộc CSDL student_db", "15"),
        ("Bảng 3.2: Danh mục 6 Security Headers bắt buộc và mục tiêu phòng vệ", "20"),
        ("Bảng 4.1: Cấu hình 4 scrape jobs định kỳ trên Prometheus Server", "21"),
        ("Bảng 4.2: Cấu trúc 5 hàng Panels giám sát trên Grafana Dashboard", "22"),
        ("Bảng 5.1: Đặc tả 3 câu truy vấn LogQL nghiệp vụ và kết quả xác thực thực nghiệm", "26"),
        ("Bảng 5.2: Tổng hợp 5 giải pháp Hardening an toàn thông tin chuyên sâu", "27")
    ]

    # Map PlantUML signatures to diagram image files
    plantuml_to_img = {
        "architecture": "architecture_10_containers.png",
        "classes": "erd_student_db.png",
        "ingress": "nginx_ingress_flow.png",
        "scrape": "prometheus_metrics_flow.png",
        "promtail": "loki_logging_flow.png"
    }

    lines = md_text.split('\n')
    start_idx = 0
    for idx, l in enumerate(lines):
        if l.strip() == '# LỜI NÓI ĐẦU':
            start_idx = idx
            break

    i = start_idx
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []
    in_toc_section = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == '---':
            i += 1
            continue

        # Code Block handling
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = '\n'.join(code_block_lines)

                c_tbl = doc.add_table(rows=1, cols=1)
                c_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_cell = c_tbl.rows[0].cells[0]
                c_cell.width = Mm(160)
                set_cell_shading(c_cell, "F9F9F9")
                set_cell_margins(c_cell, top=100, bottom=100, left=150, right=150)

                tc_border = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>\n'
                    f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
                    f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
                    f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
                    f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
                    f'</w:tcBorders>'
                )
                c_cell._tc.get_or_add_tcPr().append(tc_border)

                cp_code = c_cell.paragraphs[0]
                cp_code.paragraph_format.line_spacing = 1.15
                cp_code.paragraph_format.space_after = Pt(0)
                cp_code.paragraph_format.first_line_indent = Inches(0)

                run_c = cp_code.add_run(code_text)
                run_c.font.name = "Consolas"
                run_c.font.size = Pt(9.5)
                run_c.font.color.rgb = RGBColor(0, 0, 0)

                # Tự động chèn sơ đồ PlantUML tương ứng nếu có
                matched_diagram = None
                if "@startuml" in code_text:
                    if "frontend_net" in code_text:
                        matched_diagram = plantuml_to_img.get("architecture")
                    elif "classes (Lớp học)" in code_text or "classes" in code_text:
                        matched_diagram = plantuml_to_img.get("classes")
                    elif "Nginx Reverse Proxy (:80/:443)" in code_text:
                        matched_diagram = plantuml_to_img.get("ingress")
                    elif "Mục tiêu thu thập (Scrape Targets)" in code_text or "Scrape Engine" in code_text:
                        matched_diagram = plantuml_to_img.get("scrape")
                    elif "Promtail Log Shipper Agent" in code_text:
                        matched_diagram = plantuml_to_img.get("promtail")

                if matched_diagram:
                    d_path = os.path.join(diagrams_dir, matched_diagram)
                    if os.path.exists(d_path):
                        dp = doc.add_paragraph()
                        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        dp.paragraph_format.space_before = Pt(6)
                        dp.paragraph_format.space_after = Pt(2)
                        dp.paragraph_format.keep_with_next = True
                        dp.add_run().add_picture(d_path, width=Inches(6.0))

                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(2)
                sp.paragraph_format.line_spacing = 1.0

                code_block_lines = []
            else:
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Markdown Table handling
        if stripped.startswith('|') and '|' in stripped[1:]:
            if in_toc_section:
                i += 1
                continue
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            in_table = False
            rows_data = []
            for t_line in table_lines:
                cells = [c.strip() for c in t_line.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                rows_data.append(cells)

            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                normalized = [r + [''] * (num_cols - len(r)) for r in rows_data]

                # Kiểm tra nếu đây là KHUNG DÁN ẢNH MINH CHỨNG (Placeholder Frame)
                is_placeholder = any("[KHUNG DÁN ẢNH MINH CHỨNG" in str(r) for r in normalized)

                t_elem = doc.add_table(rows=len(normalized), cols=num_cols)
                t_elem.alignment = WD_TABLE_ALIGNMENT.CENTER

                if is_placeholder:
                    # Viền nét đứt (dashed) xám đậm để người dùng dễ nhận diện ô paste ảnh
                    set_table_borders(t_elem, color="666666", sz="12", val="dashed")
                else:
                    set_table_borders(t_elem, color="000000", sz="4", val="single")

                col_w = Mm(160 / num_cols)
                for r_idx, row_cells in enumerate(normalized):
                    doc_row = t_elem.rows[r_idx]
                    for c_idx, cell_text in enumerate(row_cells):
                        cell = doc_row.cells[c_idx]
                        cell.width = col_w
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.first_line_indent = Inches(0)
                        cp.paragraph_format.line_spacing = 1.15
                        cp.paragraph_format.space_after = Pt(2)

                        if is_placeholder:
                            if r_idx == 0:
                                set_cell_shading(cell, "EAEAEA")
                                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                add_formatted_runs(cp, cell_text, base_font_size=11.5, default_color=RGBColor(0, 0, 0))
                                for run in cp.runs:
                                    run.bold = True
                            else:
                                set_cell_shading(cell, "FAFAFA")
                                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                add_formatted_runs(cp, cell_text, base_font_size=11, default_color=RGBColor(80, 80, 80))
                                for run in cp.runs:
                                    run.italic = True
                                # Thêm khoảng trống đủ rộng để dán ảnh
                                for _ in range(4):
                                    cell.add_paragraph()
                        else:
                            if r_idx == 0:
                                set_cell_shading(cell, "F2F2F2")
                                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                add_formatted_runs(cp, cell_text, base_font_size=11, default_color=RGBColor(0, 0, 0))
                                for run in cp.runs:
                                    run.bold = True
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_shading(cell, "FFFFFF")
                                else:
                                    set_cell_shading(cell, "F9F9F9")
                                if num_cols > 2 and c_idx in (0, num_cols - 1) and len(cell_text) < 20:
                                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                else:
                                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                add_formatted_runs(cp, cell_text, base_font_size=11, default_color=RGBColor(0, 0, 0))

                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(4)
            table_lines = []

        # Images ![caption](images/path.png)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            rel_img = img_match.group(2)
            img_basename = os.path.basename(rel_img)
            full_img_path = os.path.join(images_dir, img_basename)

            if os.path.exists(full_img_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.first_line_indent = Inches(0)
                img_p.paragraph_format.space_before = Pt(8)
                img_p.paragraph_format.space_after = Pt(2)
                img_p.paragraph_format.keep_with_next = True
                img_p.add_run().add_picture(full_img_path, width=Inches(6.0))
            i += 1
            continue

        # Single Authentic Caption line: (*Hình ...* or *Bảng ...*)
        if (stripped.startswith('*Hình ') or stripped.startswith('*Bảng ')) and stripped.endswith('*'):
            cap_text = stripped[1:-1]
            cap_p = doc.add_paragraph(style='Caption')
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.first_line_indent = Inches(0)
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(8)
            r_cap = cap_p.add_run(cap_text)
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(11)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # Skip raw markdown TOC links if in TOC sections
        if stripped.startswith('- [') and '](#' in stripped:
            i += 1
            continue

        # Headings: Assign TRUE WORD STYLES & OUTLINE LEVELS
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()

            # Chapter page break
            if any(h_text.startswith(prefix) for prefix in ['LỜI NÓI ĐẦU', 'MỤC LỤC', 'DANH MỤC TỪ VIẾT TẮT', 'CHƯƠNG', 'KẾT LUẬN', 'TÀI LIỆU', 'PHỤ LỤC']):
                doc.add_page_break()

            if h_text.startswith("MỤC LỤC"):
                in_toc_section = True
            elif h_text.startswith("DANH MỤC TỪ VIẾT TẮT") or h_text.startswith("CHƯƠNG"):
                in_toc_section = False

            add_heading_with_outline(doc, h_text, level=1)

            # MỤC LỤC 1: MỤC LỤC CÁC CHƯƠNG (CHAPTER TOC)
            if "MỤC LỤC 1: MỤC LỤC CÁC CHƯƠNG" in h_text:
                p_hint = doc.add_paragraph()
                p_hint.paragraph_format.first_line_indent = Inches(0)
                p_hint.paragraph_format.space_after = Pt(6)
                r_h = p_hint.add_run("(Nhấn Ctrl + A rồi nhấn F9 trên Microsoft Word, sau đó chọn 'Update entire table' để cập nhật tự động toàn bộ mục lục)")
                r_h.font.name = "Times New Roman"
                r_h.font.italic = True
                r_h.font.size = Pt(10.5)
                r_h.font.color.rgb = RGBColor(0, 0, 0)

                p_fld_begin = doc.add_paragraph()
                p_fld_begin.paragraph_format.first_line_indent = Inches(0)
                p_fld_begin.paragraph_format.space_after = Pt(2)
                r1 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
                r2 = parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-1" \\h \\z </w:instrText></w:r>')
                r3 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
                p_fld_begin._p.append(r1)
                p_fld_begin._p.append(r2)
                p_fld_begin._p.append(r3)

                for item_title, page_num, lvl in toc_catalog:
                    if lvl == 1:
                        add_toc_entry_line(doc, item_title, page_num, level=1)

                p_fld_end = doc.add_paragraph()
                p_fld_end.paragraph_format.space_after = Pt(4)
                r4 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
                p_fld_end._p.append(r4)

            # MỤC LỤC 2: MỤC LỤC CHI TIẾT NỘI DUNG (DETAILED TOC)
            elif "MỤC LỤC 2: MỤC LỤC CHI TIẾT NỘI DUNG" in h_text:
                p_hint2 = doc.add_paragraph()
                p_hint2.paragraph_format.first_line_indent = Inches(0)
                p_hint2.paragraph_format.space_after = Pt(6)
                r_h2 = p_hint2.add_run("(Nhấn Ctrl + A rồi nhấn F9 trên Microsoft Word, sau đó chọn 'Update entire table' để cập nhật tự động toàn bộ mục lục)")
                r_h2.font.name = "Times New Roman"
                r_h2.font.italic = True
                r_h2.font.size = Pt(10.5)
                r_h2.font.color.rgb = RGBColor(0, 0, 0)

                p_fld_begin = doc.add_paragraph()
                p_fld_begin.paragraph_format.first_line_indent = Inches(0)
                p_fld_begin.paragraph_format.space_after = Pt(2)
                r1 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
                r2 = parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z </w:instrText></w:r>')
                r3 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
                p_fld_begin._p.append(r1)
                p_fld_begin._p.append(r2)
                p_fld_begin._p.append(r3)

                for item_title, page_num, lvl in toc_catalog:
                    add_toc_entry_line(doc, item_title, page_num, level=lvl)

                p_fld_end = doc.add_paragraph()
                p_fld_end.paragraph_format.space_after = Pt(4)
                r4 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
                p_fld_end._p.append(r4)

            # MỤC LỤC 3: DANH MỤC HÌNH ẢNH VÀ BẢNG BIỂU
            elif "MỤC LỤC 3: DANH MỤC HÌNH ẢNH VÀ BẢNG BIỂU" in h_text:
                # 1. Danh mục Hình ảnh
                add_heading_with_outline(doc, "Danh mục Hình ảnh (List of Figures)", level=2)
                p_hint_f = doc.add_paragraph()
                p_hint_f.paragraph_format.first_line_indent = Inches(0)
                p_hint_f.paragraph_format.space_after = Pt(4)
                r_hf = p_hint_f.add_run("(Nhấn Ctrl + A rồi nhấn F9 trên Microsoft Word để cập nhật tự động Danh mục Hình ảnh)")
                r_hf.font.name = "Times New Roman"
                r_hf.font.italic = True
                r_hf.font.size = Pt(10.5)
                r_hf.font.color.rgb = RGBColor(0, 0, 0)

                p_fld_begin = doc.add_paragraph()
                p_fld_begin.paragraph_format.first_line_indent = Inches(0)
                p_fld_begin.paragraph_format.space_after = Pt(2)
                r1 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
                r2 = parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\h \\z \\c "Hình" </w:instrText></w:r>')
                r3 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
                p_fld_begin._p.append(r1)
                p_fld_begin._p.append(r2)
                p_fld_begin._p.append(r3)

                for fig_title, page_num in figures_catalog:
                    add_toc_entry_line(doc, fig_title, page_num, level=2)

                p_fld_end = doc.add_paragraph()
                p_fld_end.paragraph_format.space_after = Pt(8)
                r4 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
                p_fld_end._p.append(r4)

                # 2. Danh mục Bảng biểu
                add_heading_with_outline(doc, "Danh mục Bảng biểu (List of Tables)", level=2)
                p_hint_t = doc.add_paragraph()
                p_hint_t.paragraph_format.first_line_indent = Inches(0)
                p_hint_t.paragraph_format.space_after = Pt(4)
                r_ht = p_hint_t.add_run("(Nhấn Ctrl + A rồi nhấn F9 trên Microsoft Word để cập nhật tự động Danh mục Bảng biểu)")
                r_ht.font.name = "Times New Roman"
                r_ht.font.italic = True
                r_ht.font.size = Pt(10.5)
                r_ht.font.color.rgb = RGBColor(0, 0, 0)

                p_fld_begin2 = doc.add_paragraph()
                p_fld_begin2.paragraph_format.first_line_indent = Inches(0)
                p_fld_begin2.paragraph_format.space_after = Pt(2)
                r1_2 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
                r2_2 = parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\h \\z \\c "Bảng" </w:instrText></w:r>')
                r3_2 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
                p_fld_begin2._p.append(r1_2)
                p_fld_begin2._p.append(r2_2)
                p_fld_begin2._p.append(r3_2)

                for tbl_title, page_num in tables_catalog:
                    add_toc_entry_line(doc, tbl_title, page_num, level=2)

                p_fld_end2 = doc.add_paragraph()
                p_fld_end2.paragraph_format.space_after = Pt(4)
                r4_2 = parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
                p_fld_end2._p.append(r4_2)

            i += 1
            continue

        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            if in_toc_section and ("Danh mục Hình ảnh" in h_text or "Danh mục Bảng biểu" in h_text):
                i += 1
                continue
            add_heading_with_outline(doc, h_text, level=2)
            i += 1
            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            add_heading_with_outline(doc, h_text, level=3)
            i += 1
            continue

        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            add_heading_with_outline(doc, h_text, level=4)
            i += 1
            continue

        if in_toc_section and stripped.startswith('*(') and stripped.endswith(')*'):
            i += 1
            continue

        # Bullet list items (- or *)
        if stripped.startswith('- ') or stripped.startswith('* '):
            list_text = stripped[2:].strip()
            lp = doc.add_paragraph()
            lp.paragraph_format.first_line_indent = Inches(-0.25)
            lp.paragraph_format.left_indent = Inches(0.5)
            lp.paragraph_format.line_spacing = 1.3
            lp.paragraph_format.space_after = Pt(3)
            lp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            r_bullet = lp.add_run("•  ")
            r_bullet.bold = True
            r_bullet.font.name = "Times New Roman"
            r_bullet.font.size = Pt(13)
            r_bullet.font.color.rgb = RGBColor(0, 0, 0)

            add_formatted_runs(lp, list_text, base_font_size=13, default_color=RGBColor(0, 0, 0))
            i += 1
            continue

        # Numbered list items
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            num_prefix = num_match.group(1) + ". "
            num_text = num_match.group(2)
            np = doc.add_paragraph()
            np.paragraph_format.first_line_indent = Inches(-0.25)
            np.paragraph_format.left_indent = Inches(0.5)
            np.paragraph_format.line_spacing = 1.3
            np.paragraph_format.space_after = Pt(3)
            np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            r_num = np.add_run(num_prefix + " ")
            r_num.bold = True
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(13)
            r_num.font.color.rgb = RGBColor(0, 0, 0)

            add_formatted_runs(np, num_text, base_font_size=13, default_color=RGBColor(0, 0, 0))
            i += 1
            continue

        # Normal Body Paragraph
        if stripped:
            bp = doc.add_paragraph()
            bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            bp.paragraph_format.first_line_indent = Inches(0.5)
            bp.paragraph_format.line_spacing = 1.3
            bp.paragraph_format.space_after = Pt(4)
            add_formatted_runs(bp, stripped, base_font_size=13, default_color=RGBColor(0, 0, 0))

        i += 1

    doc.save(docx_path)
    print(f"Generated DOCX successfully: {docx_path}")

    # Synchronize to root
    try:
        shutil.copyfile(docx_path, root_docx_path)
        print(f"Synchronized to root: {root_docx_path}")
    except PermissionError:
        backup_path = os.path.join(workspace_dir, "Bao_Cao_Trien_Khai_He_Thong_Latest.docx")
        shutil.copyfile(docx_path, backup_path)
        print(f"[NOTE] '{root_docx_path}' is currently open in Microsoft Word.")
        print(f"Saved an identical copy to: {backup_path}")
        print(f"The primary academic file is at: {docx_path}")

if __name__ == "__main__":
    build_academic_docx()
